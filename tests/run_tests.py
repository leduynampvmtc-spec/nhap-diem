"""Kiểm thử tự động (không cần pytest):  python tests/run_tests.py
Dùng dữ liệu giả lập — không chứa tên HSSV thật."""
import datetime as dt
import os
import random
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import openpyxl

from nhapdiem import controller as ctl, core, forms_pdf as F, pdfgrid, phieu_xlsx

FAILS = []


def check(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  {extra}" if extra and not cond else ""))
    if not cond:
        FAILS.append(name)


# ---------------------------------------------------------------------------
def make_input(path, a5="TÊN MÔN HỌC: CƠ KỸ THUẬT", q5="(LT: 30; TH: 15)", o5="SỐ GIỜ: 45", rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DIEM TK"
    ws["A5"], ws["O5"], ws["Q5"] = a5, o5, q5
    ws["B6"], ws["E6"], ws["O6"] = "LOP TEST 1", "KHÓA: 2025", "HỆ: CAO ĐẲNG"
    ws["AA6"] = "MÔN HỌC"
    for i, r in enumerate(rows or [], 10):
        name, tx1, tx2, dk, p, q, mssv, dob = r
        ws.cell(i, 2, name)
        ws.cell(i, 3, tx1)
        ws.cell(i, 4, tx2)
        for k, v in enumerate(dk):
            ws.cell(i, 5 + k, v)
        ws.cell(i, 16, p)
        ws.cell(i, 17, q)
        ws.cell(i, 27, mssv)
        ws.cell(i, 31, dob)
    wb.save(path)


ROWS_MH = [
    # name, tx1, tx2, dk[], P, Q, mssv, dob
    ("Nguyễn Văn An", 5.7, "-", [5.6, 8.5, 8.9], 6.7, "-", "T001", dt.datetime(2009, 10, 20)),
    ("Lê Thị Bình", 8.2, "-", [8.8, 7.7, 7.4], 4.0, "-", "T002", "04/05/2008"),
    ("Trần Văn Cường", 2.0, "-", [3.0, 3.5, 2.5], 8.0, "-", "T003", None),      # TBC < 5 -> học lại
    ("Phạm Thị Dung", 9.0, "-", [9.0, 9.5, 8.5], 9.5, "-", "T004", None),
    ("Đỗ Văn Em", 7.0, "-", [7.0, 7.0, 7.0], 2.5, "-", "T005", None),            # P < 3
    ("Hồ Thị Phương", 6.0, "-", [6.5, 6.0, 7.0], 5.0, "-", "T006", None),
]


def run():
    tmp = tempfile.mkdtemp(prefix="nhapdiem_test_")
    inp = os.path.join(tmp, "in.xlsx")
    make_input(inp, rows=ROWS_MH)

    # ---- đọc & meta ----
    meta, sts, warns = core.read_input(inp)
    check("meta: loại + tên", meta.loai == "MÔN HỌC" and meta.ten == "CƠ KỸ THUẬT", str(meta))
    check("meta: giờ + tín chỉ", (meta.gio_lt, meta.gio_th, meta.tong_gio, meta.tin_chi) == (30, 15, 45, 3))
    check("meta: lớp-khóa + tên không dấu", meta.lop_khoa == "LOP TEST 1-2025" and meta.ten_khong_dau == "CO KY THUAT")
    check("đọc đủ HSSV", len(sts) == 6)
    check("bo_dau", core.bo_dau("Đặng Thị Ngọc Ánh") == "Dang Thi Ngoc Anh")
    check("xl_round nửa ra xa 0", core.xl_round(6.35, 1) == 6.4 and core.xl_round(2.5) == 3 and core.xl_round(0.05, 1) == 0.1)

    # ---- tính điểm, đối chiếu tay ----
    core.compute_all(sts, meta)
    a, b, c, d, e, f = sts
    check("TBC KT (A) = 7.4", a.o == 7.4, a.o)
    tbc_a = (5.7 + 2 * (5.6 + 8.5 + 8.9)) / 7
    check("Tổng kết (A) = round(6.7*.6+tbc*.4)", a.s == core.xl_round(6.7 * 0.6 + tbc_a * 0.4, 1) and a.x == "Đạt", (a.s, a.x))
    check("A: Q = '-' (môn học)", a.q == "-" and a.r == 6.7)
    check("B thi lại (P=4.0, R<5)", b.x == "Thi lại" and b.y == "Điểm thi lần thứ nhất < 5", (b.x, b.y))
    check("C học lại (O<5)", c.x == "Học lại" and c.p == "-" and c.r == "-" and c.s == "-" and c.y == "Nghỉ học quá số buổi quy định", (c.x, c.y, c.s))
    check("D đạt", d.x == "Đạt")
    check("E thi lại (P<3)", e.x == "Thi lại")
    check("F: P=5.0 đạt", f.x == "Đạt", (f.x, f.s))

    # tín chỉ > 3 => dùng KT thường xuyên 2
    m2 = core.Meta(loai="MÔN HỌC", ten="X", gio_lt=45, gio_th=30, tin_chi=4)
    s2 = core.Student(row=10, stt=1, name="T", tx1=6.0, tx2=8.0, dk=[7.0] + [None] * 9)
    core.compute_all([s2], m2)
    check("4 tín chỉ: dùng cả TX1, TX2", s2.o == core.xl_round((6 + 8 + 14) / 4, 1) and s2.tx2 == 8.0, s2.o)

    # mô đun có TH: R, S theo trọng số giờ
    mm = core.Meta(loai="MÔ ĐUN", ten="Y", gio_lt=15, gio_th=60, tin_chi=3)
    s3 = core.Student(row=10, stt=1, name="M", tx1=8.0, tx2="-", dk=[8.0] + [None] * 9, p=6.0, q=8.0)
    core.compute_all([s3], mm)
    wl, wt = 15 / 15, 60 / 30
    r_exp = core.xl_round((6 * wl + 8 * wt) / (wl + wt), 1)
    s_exp = core.xl_round((6 * wl + 8 * wt) / (wl + wt) * 0.6 + 8.0 * 0.4, 1)
    check("Mô đun: R, S theo giờ LT/TH", s3.r == r_exp and s3.s == s_exp and s3.x == "Đạt", (s3.r, s3.s, s3.x))
    s4 = core.Student(row=10, stt=1, name="M2", tx1=8.0, tx2="-", dk=[8.0] + [None] * 9, p=7.0, q=2.5)
    core.compute_all([s4], mm)
    check("Mô đun: TH < 3 => thi lại", s4.x == "Thi lại", (s4.x, s4.y))

    # ---- Xét ĐK ----
    ok_c, note_c = core.xet_dieu_kien(c)
    check("Xét ĐK: TBC<5 => không đủ + Học lại", (ok_c, note_c) == (False, "Học lại"))
    check("Xét ĐK: đủ", core.xet_dieu_kien(a) == (True, ""))
    a.vang_tong = 0.25
    check("Xét ĐK 2022: vắng > 20% => không đủ", core.xet_dieu_kien(a, "2022")[0] is False)
    a.vang_tong, a.vang_lt, a.vang_th = 0.1, 0.25, 0.0
    check("Xét ĐK 2021: vắng LT 25% đủ", core.xet_dieu_kien(a, "2021")[0] is True)
    a.vang_lt, a.vang_th = 0.35, 0.0
    check("Xét ĐK 2021: vắng LT 35% không đủ", core.xet_dieu_kien(a, "2021")[0] is False)
    a.vang_tong = a.vang_lt = a.vang_th = None

    # ---- Thi lại ----
    tl = core.danh_sach_thi_lai(sts, meta, "Thi trắc nghiệm trên máy tính", "Thi thực hành tại xưởng TBĐ")
    # VBA so sánh chuỗi nhị phân, giảm dần: "Đ" (U+0110) lớn hơn "L" nên "Đỗ Văn Em" đứng trước
    check("Thi lại: 2 HSSV, tên giảm dần (so sánh nhị phân như VBA)", [x["name"] for x in tl] == ["Đỗ Văn Em", "Lê Thị Bình"], [x["name"] for x in tl])
    check("Thi lại: ghi chú LT", all(x["lt"] and not x["th"] and x["ghichu"] == "Thi trắc nghiệm trên máy tính" for x in tl))
    mm_students = [core.Student(row=10, stt=1, name="A1", tx1=8, tx2="-", dk=[8] + [None] * 9, p=3.0, q=3.0),
                   core.Student(row=11, stt=2, name="B1", tx1=8, tx2="-", dk=[8] + [None] * 9, p=3.0, q=9.0),
                   core.Student(row=12, stt=3, name="C1", tx1=8, tx2="-", dk=[8] + [None] * 9, p=9.0, q=4.0)]
    mm2 = core.Meta(loai="MÔ ĐUN", ten="Y", gio_lt=15, gio_th=60, tin_chi=3)
    core.compute_all(mm_students, mm2)
    for s_ in mm_students:
        s_.x = "Thi lại"
    t2 = core.danh_sach_thi_lai(mm_students, mm2, "LT-TXT", "TH-TXT")
    by = {x["name"]: x for x in t2}
    check("Mô đun (TH>LT): cả 2 / chỉ LT / chỉ TH",
          by["A1"]["lt"] and by["A1"]["th"] and by["A1"]["ghichu"] == "LT-TXT và TH-TXT"
          and by["B1"]["lt"] and not by["B1"]["th"] and by["B1"]["ghichu"] == "LT-TXT"
          and by["C1"]["th"] and not by["C1"]["lt"] and by["C1"]["ghichu"] == "TH-TXT", str(t2))

    # ---- LT ----
    lt1 = core.bang_diem_lt(sts, 1)
    check("LT lần 1: 6 dòng, ngày sinh dd/mm/yyyy", len(lt1) == 6 and lt1[0]["dob"] == "20/10/2009" and lt1[0]["diem"] == "6.7"
          and lt1[1]["dob"] == "04/05/2008", lt1[:2])
    sts[1].t = 5.5
    lt2 = core.bang_diem_lt(sts, 2)
    check("LT lần 2: chỉ HSSV có LT lần 1 < 5", [x["name"] for x in lt2] == ["Lê Thị Bình", "Đỗ Văn Em"] and lt2[0]["diem"] == "5.5" and lt2[1]["diem"] == "", lt2)

    # ---- Xếp thứ tự ----
    xt = core.xep_thu_tu(sts)
    check("XTT: ≥7 giảm dần", [n for n, _ in xt["lon7"]] == ["Phạm Thị Dung", "Nguyễn Văn An"], xt["lon7"])
    check("XTT: '-' đứng đầu bảng <5", xt["nho5"][0] == ("Trần Văn Cường", "-"), xt["nho5"])
    check("XTT: top3 / bottom3", len(xt["top3"]) == 3 and xt["top3"][0][0] == "Phạm Thị Dung"
          and xt["bottom3"][0] == ("Trần Văn Cường", "-") and len(xt["bottom3"]) == 4, xt)

    # ---- tách điểm phiếu TH ----
    rng = random.Random(1)
    bad = 0
    for t10 in range(0, 101):
        sc = t10 / 10
        tcs = phieu_xlsx.split_score(sc, rng=rng)
        frac = [x for x in tcs if round(x * 10) % 5 != 0]
        if abs(sum(tcs) - sc) > 1e-9 or any(x < 0 or x > 1 for x in tcs) or len(tcs) != 10 or len(frac) > 1:
            bad += 1
    check("Tách điểm TC: mọi điểm 0.0–10.0", bad == 0, f"{bad} lỗi")

    # ---- điểm danh (bố cục thực tế: X / P / KP, trọng số theo giờ) ----
    att = os.path.join(tmp, "att.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Khac"
    w = wb.create_sheet("CKT")
    names = ["Nguyễn Văn An", "LÊ THỊ BÌNH", "Trần  Văn Cường"]
    for i, nm in enumerate(names, 12):
        w.cell(i, 1, i - 11)
        w.cell(i, 2, nm)
    kinds = ["LT"] * 8 + ["TH"] * 2
    for j, k in enumerate(kinds):
        w.cell(10, 4 + j, k)
        w.cell(11, 4 + j, 4)              # mỗi buổi 4 giờ -> 40 giờ
        for r in (12, 13, 14):
            w.cell(r, 4 + j, "X")
    w.cell(11, 4 + 10 + 3, "x")           # cột phụ bên phải: bị bỏ qua
    w.cell(10, 4 + 10 + 2, "Số buổi vắng")
    for j in (0, 1, 2):                   # An vắng 3 buổi LT (3/10 buổi = 30%)
        w.cell(12, 4 + j, "P")
    w.cell(13, 4 + 9, "KP")               # Bình vắng 1 buổi TH (1/10 buổi = 10%)
    wb.save(att)
    a_map = core.load_attendance(att, names=names)
    check("Điểm danh: tổng/LT/TH theo số buổi", a_map["NGUYEN VAN AN"] == (0.3, 0.3, 0.0) and a_map["LE THI BINH"] == (0.1, 0.0, 0.1), a_map)
    miss = core.apply_attendance(sts, a_map)
    check("Điểm danh: gán theo tên không dấu, báo HSSV không khớp", sts[0].vang_tong == 0.3 and len(miss) == 3, miss)
    check("Xét ĐK với điểm danh (An vắng 30%)", core.xet_dieu_kien(sts[0])[0] is False)
    for s_ in sts:
        s_.vang_tong = s_.vang_lt = s_.vang_th = None

    # ---- đăng nhập ----
    from nhapdiem import auth
    check("Đăng nhập đúng", auth.check_login("leduynam", "06241984") and auth.check_login(" LeDuyNam ", "06241984"))
    check("Đăng nhập sai bị từ chối", not auth.check_login("leduynam", "x") and not auth.check_login("abc", "06241984") and not auth.check_login("", ""))
    check("Logo có trong gói", all(os.path.isfile(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "nhapdiem", "assets", n)) for n in ("logo_240.png", "logo_72.png", "icon.ico")))

    # ---- xuất file ----
    out = os.path.join(tmp, "out")
    sess = ctl.load_session(inp)
    p1 = ctl.export_dieu_kien(sess, out, F.DKOptions(ma_lhp="010100046101", dia_diem="Xưởng TBĐ"))
    p2 = ctl.export_thi_lai(sess, out, F.TLOptions(ma_lhp="010100046101", dia_diem_th="xưởng thực hành"))
    try:
        ctl.export_diem_lt(sess, out, F.LTOptions(lan_thi=1, ca_thi="Ca 1"))
        blocked = False
    except ctl.AppError as e:
        blocked = "MÔ ĐUN" in str(e)
    check("LT: MÔN HỌC bị từ chối, chỉ xuất cho MÔ ĐUN", blocked)
    inp_md = os.path.join(tmp, "in_modun.xlsx")
    make_input(inp_md, a5="TÊN MÔ ĐUN: SỬA CHỮA MÁY NÉN KHÍ", rows=ROWS_MH)
    sess_md = ctl.load_session(inp_md)
    p3 = ctl.export_diem_lt(sess_md, out, F.LTOptions(lan_thi=1, ca_thi="Ca 1 (1/1/2026)"))
    p4 = ctl.export_xep_thu_tu(sess, out, dt.date(2026, 10, 21))
    p5, n5 = ctl.export_phieu(sess, out, 2, dt.date(2026, 10, 21), random.Random(5))
    for p in (p1, p2, p3, p4):
        with open(p, "rb") as fh:
            head = fh.read(5)
        check(f"PDF hợp lệ: {os.path.basename(p)}", head == b"%PDF-" and os.path.getsize(p) > 2000)
    check("Tên file giống VBA", os.path.basename(p1).endswith("_DSDT_CO KY THUAT_LOP TEST 1-2025.pdf"), os.path.basename(p1))
    check("Phiếu xlsx: đủ HSSV", n5 == 6)
    wsx = openpyxl.load_workbook(p5)["DIEM TC"]
    ok = True
    for r in range(2, 8):
        score = wsx.cell(r, 3).value
        tcs = [wsx.cell(r, c).value for c in range(5, 15)]
        ok &= abs(sum(tcs) - score) < 1e-9 and abs(wsx.cell(r, 15).value - score) < 1e-9
    check("Phiếu xlsx: Σ TC1..TC10 = điểm KT (bài 2)", ok)
    try:
        ctl.export_phieu(sess, out, 5, dt.date.today())
        check("Phiếu bài chưa có điểm báo lỗi", False)
    except ctl.AppError:
        check("Phiếu bài chưa có điểm báo lỗi", True)


    # ---- xlsx + footer cho 4 file xuất ----
    from nhapdiem.footer import FooterOptions
    ftc = FooterOptions(codes={"dk": "BM/CL/05.01", "tl": "BM/CL/05.02", "lt": "BM/CL/05.04a", "tk": ""},
                         lbh={"dk": "07", "tl": "08", "lt": "09", "tk": "10"},
                         nhl={"dk": "01/09/2026", "tl": "02/09/2026", "lt": "03/09/2026", "tk": "04/09/2026"})
    dd = dt.date(2026, 10, 21)
    out_x = os.path.join(tmp, "out_x")
    xs = {
        "dk": (ctl.export_dieu_kien(sess, out_x, F.DKOptions(ma_lhp="1"), "xlsx", ftc), "DK DT"),
        "tl": (ctl.export_thi_lai(sess, out_x, F.TLOptions(ma_lhp="1"), "xlsx", ftc), "TL"),
        "lt": (ctl.export_diem_lt(sess_md, out_x, F.LTOptions(ca_thi="Ca 1"), "xlsx", ftc), "LT"),
        "tk": (ctl.export_xep_thu_tu(sess, out_x, dd, "xlsx", ftc), "TK"),
    }
    for k, (pth, sheet) in xs.items():
        wsf = openpyxl.load_workbook(pth)[sheet]
        check(f"xlsx {k}: đuôi .xlsx, font TNR", pth.endswith(".xlsx") and wsf["A1"].font.name == "Times New Roman" or k == "tk")
        if k == "tl":
            check("Footer: LBH/NHL khác nhau theo file", ftc.center("tl") == "LBH: 08      NHL: 02/09/2026" and ftc.center("dk") != ftc.center("tl"))
        check(f"xlsx {k}: footer riêng LBH/NHL + Trang &P/&N", wsf.oddFooter.center.text == ftc.center(k)
              and wsf.oddFooter.right.text == "Trang &P/&N" and wsf.oddFooter.left.text == ftc.codes[k])
    pf = ctl.export_dieu_kien(sess, out_x, F.DKOptions(ma_lhp="1"), "pdf", ftc)
    try:
        tf = subprocess.run(["pdftotext", "-layout", pf, "-"], capture_output=True, text=True, timeout=30).stdout
        check("PDF: footer mỗi trang", "BM/CL/05.01" in tf and "LBH: 07" in tf and "NHL: 01/09/2026" in tf and "Trang 1/" in tf)
    except FileNotFoundError:
        pass


    # ---- Tab 6: điền Form nhập điểm hệ thống ----
    fx = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
    sm = [x for x in sess_md.students if x.mssv][:3]
    for fn, tag in (("mau_qua_trinh.xlsx", "qt"), ("mau_ket_thuc.xlsx", "kt")):
        wbt = openpyxl.load_workbook(os.path.join(fx, fn))
        wst = wbt.active
        col = 2
        for i, stu in enumerate(sm):
            wst.cell(10 + i, col).value = stu.mssv
        wbt.save(os.path.join(tmp, f"tpl_{tag}.xlsx"))
    if sm:
        pq, wq = ctl.export_nhap_qua_trinh(sess_md, os.path.join(tmp, "tpl_qt.xlsx"), out_x, dd)
        oq = openpyxl.load_workbook(pq).active
        check("Tab6 quá trình: điền đúng TBTK/tổng kết theo MSSV",
              oq.cell(10, 13).value == (round(sm[0].o, 1) if core.is_num(sm[0].o) else None), str(wq)[:120])
        pk, wk = ctl.export_nhap_ket_thuc(sess_md, os.path.join(tmp, "tpl_kt.xlsx"), out_x, dd, 1)
        ok_ = openpyxl.load_workbook(pk).active
        check("Tab6 kết thúc mô đun: điền LT/TH", ok_.cell(10, 7).value == (round(sm[0].p, 1) if core.is_num(sm[0].p) else None))
    try:
        ctl.export_nhap_qua_trinh(sess_md, os.path.join(tmp, "x.xls"), out_x, dd)
        check("Tab6: thiếu file mẫu báo lỗi", False)
    except ctl.AppError:
        check("Tab6: thiếu file mẫu báo lỗi", True)

    # kiểm tra chữ trong PDF (nếu có pdftotext)
    try:
        txt = subprocess.run(["pdftotext", "-layout", p1, "-"], capture_output=True, text=True, timeout=30).stdout
        check("PDF ĐK: có tiếng Việt, mã LHP, địa điểm", "DANH SÁCH XÉT ĐIỀU KIỆN" in txt and "010100046101" in txt and "Xưởng TBĐ" in txt)
        txt2 = subprocess.run(["pdftotext", "-layout", p2, "-"], capture_output=True, text=True, timeout=30).stdout
        check("PDF TL: chỉ có HSSV thi lại", "Lê Thị Bình" in txt2 and "Đỗ Văn Em" in txt2 and "Nguyễn Văn An" not in txt2)
    except FileNotFoundError:
        print("SKIP kiểm tra chữ trong PDF (không có pdftotext)")

    # nhiều trang: 70 HSSV -> chữ ký không bị tách riêng
    many = [(f"Học Viên {i:02d}", 7.0, "-", [7.0, 7.5, 8.0], 7.0, "-", f"T{i:03d}", None) for i in range(70)]
    inp2 = os.path.join(tmp, "many.xlsx")
    make_input(inp2, rows=many)
    s70 = ctl.load_session(inp2)
    pp = ctl.export_dieu_kien(s70, out, F.DKOptions(ma_lhp="1"))
    pages = subprocess.run(["pdfinfo", pp], capture_output=True, text=True).stdout if subprocess.run(["which", "pdfinfo"], capture_output=True).returncode == 0 else ""
    if pages:
        n = int([l for l in pages.splitlines() if l.startswith("Pages:")][0].split()[1])
        txt = subprocess.run(["pdftotext", "-layout", "-f", str(n), "-l", str(n), pp, "-"], capture_output=True, text=True).stdout
        check("70 HSSV: trang cuối có cả dòng dữ liệu và chữ ký", "GIÁO VIÊN GIẢNG DẠY" in txt and "Học Viên" in txt, f"{n} trang")

    # lỗi thân thiện
    for bad_path in ("", os.path.join(tmp, "khong_co.xlsx")):
        try:
            ctl.load_session(bad_path)
            check("Báo lỗi file không tồn tại", False)
        except ctl.AppError:
            check("Báo lỗi file không tồn tại", True)
    check("Ngày lập sai báo lỗi", _raises(lambda: ctl.parse_date("32/13/2026")))
    check("Ngày lập đúng", ctl.parse_date("21/10/2026") == dt.date(2026, 10, 21))

    # file thật (chỉ chạy khi có)
    real = "/mnt/user-data/uploads/Form_nhap_diem.xlsx"
    if os.path.exists(real):
        sr = ctl.load_session(real)
        check("File thật: 28 HSSV, O khớp công thức Excel",
              len(sr.students) == 28 and _o_matches_excel(real, sr))


def _raises(fn):
    try:
        fn()
    except ctl.AppError:
        return True
    return False


def _o_matches_excel(path, sess):
    ws = openpyxl.load_workbook(path, data_only=True)["DIEM TK"]
    for s in sess.students:
        cached = ws.cell(s.row, 15).value
        if cached is None or abs(float(cached) - s.o) > 1e-9:
            return False
    return True


def gui_smoke():
    """Chạy toàn bộ code GUI với tkinter giả (sandbox không có Tk) để bắt lỗi tên/thuộc tính."""
    import types
    if "tkinter" in sys.modules and not isinstance(sys.modules["tkinter"], types.ModuleType):
        return
    try:
        import tkinter  # noqa
        print("SKIP gui_smoke (có tkinter thật)")
        return
    except ImportError:
        pass

    class W:
        def __init__(self, *a, **k):
            pass

        def __getattr__(self, n):
            return lambda *a, **k: W()

    class Var:
        def __init__(self, value="", **k):
            self.v = value

        def get(self):
            return self.v

        def set(self, v):
            self.v = v

    tk = types.ModuleType("tkinter")
    tk.Tk, tk.Text = W, W
    tk.StringVar = tk.BooleanVar = tk.IntVar = Var
    ttk = types.ModuleType("tkinter.ttk")
    for n in ("Label", "Entry", "Button", "Frame", "LabelFrame", "Notebook", "Scrollbar", "Checkbutton", "Radiobutton"):
        setattr(ttk, n, W)
    asked = []
    fd = types.ModuleType("tkinter.filedialog")
    fd.askopenfilename = lambda **k: ""
    fd.askdirectory = lambda **k: ""
    mb = types.ModuleType("tkinter.messagebox")
    mb.askyesno = lambda *a, **k: (asked.append(a) or False)
    mb.showwarning = mb.showerror = mb.showinfo = lambda *a, **k: asked.append(a)
    tk.ttk, tk.filedialog, tk.messagebox = ttk, fd, mb
    sys.modules.update({"tkinter": tk, "tkinter.ttk": ttk, "tkinter.filedialog": fd, "tkinter.messagebox": mb})
    for m in [m for m in sys.modules if m.startswith("nhapdiem.gui")]:
        del sys.modules[m]
    from nhapdiem import gui
    gui.CONFIG_PATH = os.path.join(tempfile.mkdtemp(), "cfg.json")
    app = gui.App()
    tmp = tempfile.mkdtemp()
    inp = os.path.join(tmp, "in.xlsx")
    make_input(inp, a5="TÊN MÔ ĐUN: SỬA CHỮA MÁY NÉN KHÍ", rows=ROWS_MH)
    app.v_in.set(inp)
    app.v_out.set(os.path.join(tmp, "o"))
    app.v_malhp.set("123")
    app.v_ca.set("Ca 4 - Ca 7 (23/1/2026)")
    for fn in (app.do_dk, app.do_tl, app.do_lt, app.do_xtt, app.do_phieu):
        fn()
    app.v_lan.set(2)
    app.do_lt()
    app.v_bai.set(4)
    app.do_phieu()          # bài 4 trống -> báo lỗi thân thiện, không văng
    app._on_close()
    made = sorted(os.listdir(os.path.join(tmp, "o")))
    check("GUI (tk giả): các nút xuất tạo đủ file", len(made) >= 5, str(made))


def web_smoke():
    """Kiểm thử giao diện web cục bộ qua HTTP thật (đăng nhập, tải tệp, xuất file, chặn truy cập lạ)."""
    import http.client
    import json
    import threading
    from urllib.parse import quote
    from nhapdiem import webapp
    tmp = tempfile.mkdtemp()
    webapp.CONFIG_PATH = os.path.join(tmp, "webcfg.json")
    srv = webapp.make_server()
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    cookie = {}

    def call(method, path, body=None, raw=False, headers=None):
        c = http.client.HTTPConnection("127.0.0.1", port, timeout=20)
        h = dict(headers or {})
        if cookie.get("v"):
            h["Cookie"] = "nd_session=" + cookie["v"]
        data = None
        if body is not None:
            data = body if raw else json.dumps(body).encode()
            if not raw:
                h["Content-Type"] = "application/json"
        c.request(method, path, body=data, headers=h)
        r = c.getresponse()
        b = r.read()
        sc = r.getheader("Set-Cookie")
        if sc and sc.startswith("nd_session="):
            cookie["v"] = "" if "Max-Age=0" in sc else sc.split(";")[0].split("=", 1)[1]
        c.close()
        try:
            return r.status, json.loads(b.decode())
        except Exception:
            return r.status, b

    check("Web: trang chủ có tiêu đề", "Xuất biểu mẫu".encode() in call("GET", "/")[1])
    check("Web: chưa đăng nhập bị chặn", call("GET", "/api/config")[0] == 401 and call("POST", "/api/load", {})[0] == 401)
    check("Web: chặn Host lạ (DNS rebinding)", call("GET", "/", headers={"Host": "evil.com"})[0] == 403)
    check("Web: chặn Origin lạ", call("POST", "/api/login", {"user": "leduynam", "password": "06241984"},
                                      headers={"Origin": "http://evil.com"})[0] == 403)
    check("Web: sai mật khẩu bị từ chối", call("POST", "/api/login", {"user": "leduynam", "password": "x"})[0] == 401)
    st, d = call("POST", "/api/login", {"user": "leduynam", "password": "06241984"})
    check("Web: đăng nhập đúng", st == 200 and d.get("ok") and cookie.get("v"))
    inp = os.path.join(tmp, "in.xlsx")
    make_input(inp, rows=ROWS_MH)
    with open(inp, "rb") as fh:
        st, d = call("POST", "/api/upload?kind=input&name=" + quote("Form nhập điểm.xlsx"), fh.read(), raw=True)
    check("Web: tải tệp nhập điểm", st == 200 and d.get("ok"))
    st, d = call("POST", "/api/upload?kind=input&name=evil.exe", b"MZ", raw=True)
    check("Web: từ chối tệp không phải Excel", st == 400)
    st, d = call("POST", "/api/load", {"drop": False})
    check("Web: đọc dữ liệu", st == 200 and d["meta"]["n"] == len(ROWS_MH) and d["meta"]["la_mo_dun"] is False, str(d)[:200])
    out = os.path.join(tmp, "o")
    base = {"out_dir": out, "ngay": "21/10/2026", "ma_lhp": "123", "tbm": "CƠ KHÍ", "dia_diem": "Phòng máy",
            "quy_che": "2022", "ht_lt": "Thi trắc nghiệm trên máy tính", "dd_th": "xưởng TBĐ", "tn_may": True, "bai": 2}
    for kind in ("dk", "tl", "xtt", "phieu"):
        st, d = call("POST", "/api/export/" + kind, base)
        ok = st == 200 and d.get("ok") and os.path.isfile(d["path"])
        check(f"Web: xuất {kind}", ok, str(d)[:200])
        if ok:
            st2, body = call("GET", "/download/" + d["id"])
            check(f"Web: tải về {kind}", st2 == 200 and len(body) > 1000)
    st, d = call("POST", "/api/export/lt", dict(base, ca_thi="Ca 1"))
    check("Web: LT bị chặn với MÔN HỌC", st == 400 and "MÔ ĐUN" in d.get("error", ""))
    check("Web: đăng xuất", call("POST", "/api/logout", {})[0] == 200 and call("GET", "/api/config")[0] == 401)
    srv.shutdown()


if __name__ == "__main__":
    pdfgrid.ensure_fonts()
    run()
    gui_smoke()
    web_smoke()
    print("\n" + ("TẤT CẢ ĐỀU ĐẠT" if not FAILS else f"CÓ {len(FAILS)} LỖI: {FAILS}"))
    sys.exit(1 if FAILS else 0)
