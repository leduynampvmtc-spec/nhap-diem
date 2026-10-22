"""
Điền điểm vào 2 file Form nhập điểm của hệ thống nhà trường (mẫu [ASCIMP]):
  - Form nhập quá trình      : DiemHeSo11/12 (KT thường xuyên), DiemHeSo21..25 (KT định kỳ 1-5),
                               DiemTBThuongKy (TBC KT), DuocDuThiKetThuc ("V"), DiemTongKet
  - Form nhập thi kết thúc mô đun : DiemThiKN1 (LT), DiemThiKN2 (TH)
Các cột được nhận diện bằng "mã cột" nằm ở hàng 6 của file mẫu, HSSV khớp theo Mã sinh viên
(không có thì khớp theo họ tên không dấu). Giữ nguyên định dạng, ô gộp, công thức tổng của file mẫu.
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional, Tuple

import openpyxl

from . import core
from .core import Meta, Student, bo_dau, is_num, xet_dieu_kien


class TemplateError(Exception):
    pass


def _key_name(s: object) -> str:
    return " ".join(bo_dau(str(s or "")).upper().split())


def _scan(ws) -> Tuple[Dict[str, int], int, int, int, int]:
    """(mã cột -> số cột, hàng mã, cột MSSV, cột họ đệm, cột tên)."""
    codes: Dict[str, int] = {}
    code_row = 0
    for r in range(1, 21):
        row = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and re.match(r"^(Diem|DuocDu)[A-Za-z0-9]+$", v.strip()):
                row[v.strip()] = c
        if row:
            codes, code_row = row, r
            break
    if not codes:
        raise TemplateError("Không thấy hàng mã cột (DiemHeSo11, DiemThiKN1…) ở 20 hàng đầu — "
                            "đây không phải file Form nhập điểm của hệ thống.")
    c_mssv = c_ho = c_ten = 0
    for r in range(1, 21):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            t = _key_name(v)
            if t == "MA SINH VIEN":
                c_mssv = c
            elif t == "HO DEM":
                c_ho = c
            elif t == "TEN":
                c_ten = c
    if not c_mssv:
        raise TemplateError("Không tìm thấy cột \"Mã sinh viên\" trong file mẫu.")
    return codes, code_row, c_mssv, c_ho, c_ten


def _template_rows(ws, code_row: int, c_mssv: int) -> List[int]:
    rows = []
    for r in range(code_row + 1, ws.max_row + 1):
        a, m = ws.cell(r, 1).value, ws.cell(r, c_mssv).value
        if isinstance(a, (int, float)) and m not in (None, ""):
            rows.append(r)
    return rows


def _match(ws, rows, c_mssv, c_ho, c_ten, students: List[Student]):
    by_ms = {s.mssv.strip().upper(): s for s in students if s.mssv}
    by_name = {_key_name(s.name): s for s in students}
    pairs, miss, used = [], [], set()
    for r in rows:
        ms = str(ws.cell(r, c_mssv).value or "").strip().upper()
        st = by_ms.get(ms)
        if st is None and c_ten:
            full = f"{ws.cell(r, c_ho).value or ''} {ws.cell(r, c_ten).value or ''}" if c_ho else str(ws.cell(r, c_ten).value or "")
            st = by_name.get(_key_name(full))
        if st is None:
            miss.append(f"{ms} {ws.cell(r, c_ho).value or ''} {ws.cell(r, c_ten).value or ''}".strip())
        else:
            used.add(id(st))
            pairs.append((r, st))
    extra = [s.name for s in students if id(s) not in used]
    return pairs, miss, extra


def _num(v) -> Optional[float]:
    return round(float(v), 1) if is_num(v) else None


def _open(path: str):
    if str(path).lower().endswith(".xls"):
        raise TemplateError("File mẫu đang là .xls (Excel cũ). Hãy mở bằng Excel → Save As → "
                            "\"Excel Workbook (.xlsx)\" rồi chọn lại.")
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:  # noqa: BLE001
        raise TemplateError(f"Không đọc được file mẫu: {e}")
    return wb, wb.worksheets[0]


def _report(pairs, miss, extra, label) -> List[str]:
    w = [f"{label}: đã điền {len(pairs)} HSSV."]
    if miss:
        w.append(f"{len(miss)} HSSV trong file mẫu không có trong Form nhập điểm (để trống): " + "; ".join(miss[:8])
                 + (" …" if len(miss) > 8 else ""))
    if extra:
        w.append(f"{len(extra)} HSSV trong Form nhập điểm không có trong file mẫu (bỏ qua): " + "; ".join(extra[:8])
                 + (" …" if len(extra) > 8 else ""))
    return w


def fill_qua_trinh(tpl: str, out: str, meta: Meta, students: List[Student], quy_che: str = "2022") -> List[str]:
    wb, ws = _open(tpl)
    codes, code_row, c_ms, c_ho, c_ten = _scan(ws)
    rows = _template_rows(ws, code_row, c_ms)
    if not rows:
        raise TemplateError("File mẫu không có dòng HSSV nào.")
    pairs, miss, extra = _match(ws, rows, c_ms, c_ho, c_ten, students)
    for r, st in pairs:
        for code, col in codes.items():
            m = re.match(r"^DiemHeSo([12])(\d)$", code)
            val = None
            if m:
                i = int(m.group(2)) - 1
                src = [st.tx1, st.tx2] if m.group(1) == "1" else st.dk
                val = _num(src[i]) if 0 <= i < len(src) else None
            elif code == "DiemTBThuongKy":
                val = _num(st.o)
            elif code == "DuocDuThiKetThuc":
                val = "V" if xet_dieu_kien(st, quy_che)[0] else None
            elif code == "DiemTongKet":
                val = _num(st.s)
            else:
                continue
            ws.cell(r, col).value = val
    wb.save(out)
    return _report(pairs, miss, extra, "Nhập quá trình")


def fill_ket_thuc(tpl: str, out: str, meta: Meta, students: List[Student], lan: int = 1) -> List[str]:
    wb, ws = _open(tpl)
    codes, code_row, c_ms, c_ho, c_ten = _scan(ws)
    if "DiemThiKN1" not in codes and "DiemThiKN2" not in codes:
        raise TemplateError("File mẫu không có cột DiemThiKN1 / DiemThiKN2 (điểm thi LT / TH).")
    rows = _template_rows(ws, code_row, c_ms)
    if not rows:
        raise TemplateError("File mẫu không có dòng HSSV nào.")
    pairs, miss, extra = _match(ws, rows, c_ms, c_ho, c_ten, students)
    for r, st in pairs:
        lt, th = (st.p, st.q) if lan == 1 else (st.t, st.u)
        if "DiemThiKN1" in codes:
            ws.cell(r, codes["DiemThiKN1"]).value = _num(lt)
        if "DiemThiKN2" in codes:
            ws.cell(r, codes["DiemThiKN2"]).value = _num(th)
    wb.save(out)
    w = _report(pairs, miss, extra, f"Thi kết thúc mô đun lần {lan}")
    if not meta.la_mo_dun:
        w.append(f"Lưu ý: Form nhập điểm đang đọc là \"{meta.loai or 'không rõ'}\", không phải MÔ ĐUN.")
    n_empty = sum(1 for _, s in pairs if not (is_num((s.p, s.t)[lan - 1]) or is_num((s.q, s.u)[lan - 1])))
    if n_empty:
        w.append(f"{n_empty} HSSV chưa có điểm thi lần {lan} (để trống).")
    return w
