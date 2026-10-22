"""
Phiếu đánh giá kết quả thực hành (sheet "DIEM TC" trong Form Code VBA.xlsm) — xuất .xlsx.

Port từ TC.bas:
  Bai1..Bai5  -> lấy điểm KT định kỳ 1..5 (cột E..I của sheet DIEM TK)
  Diem_TC     -> tách điểm thành 10 tiêu chí TC1..TC10, mỗi tiêu chí 0 / 0.5 / 1,
                 tổng 10 tiêu chí đúng bằng điểm KT.

Lưu ý: VBA gốc lặp ngẫu nhiên cho tới khi tổng = điểm (chỉ thoát được khi điểm là bội của 0.5,
còn điểm kiểu 5.7 sẽ bị treo). Bản Python xử lý trọn vẹn: phần lẻ (0.1–0.4) được cộng vào đúng một
tiêu chí chưa đạt tối đa — đúng tinh thần bước "chỉnh tiêu chí cuối" trong VBA.
"""
from __future__ import annotations

import random
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .core import DASH, Meta, Student, is_num

N_CRITERIA = 10


def split_score(target: float, n: int = N_CRITERIA, rng: Optional[random.Random] = None) -> List[float]:
    """Tách `target` (0..10, 1 chữ số thập phân) thành n tiêu chí, tổng đúng bằng target."""
    rng = rng or random
    T = int(round(target * 10))
    T = max(0, min(T, 10 * n))
    units_total = (T // 5)           # số "nửa điểm"
    rem = T - units_total * 5        # 0..4 phần mười
    slots = [i for i in range(n) for _ in range(2)]   # mỗi tiêu chí tối đa 2 nửa điểm
    rng.shuffle(slots)
    units = [0] * n
    for i in slots[:units_total]:
        units[i] += 1
    tenths = [u * 5 for u in units]
    if rem:
        cand = [i for i in range(n) if tenths[i] + rem <= 10]
        tenths[rng.choice(cand)] += rem
    assert sum(tenths) == T
    return [t / 10 for t in tenths]


def xuat_phieu_xlsx(path: str, meta: Meta, students: List[Student], bai: int,
                    rng: Optional[random.Random] = None) -> int:
    """Ghi phiếu cho bài `bai` (1..5). Trả về số HSSV có điểm."""
    if not 1 <= bai <= 5:
        raise ValueError("Chỉ hỗ trợ bài 1..5 của điểm kiểm tra định kỳ.")
    wb = Workbook()
    ws = wb.active
    ws.title = "DIEM TC"

    headers = ["Stt", "Họ và tên", "Điểm KT", "Tính"] + [f"TC{i}" for i in range(1, N_CRITERIA + 1)] + ["Tính", "Đề"]
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font_b = Font(name="Times New Roman", size=13, bold=True)
    font_n = Font(name="Times New Roman", size=13)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.font, c.alignment, c.border = font_b, center, border
    ws.row_dimensions[1].height = 21

    count = 0
    for i, st in enumerate(students, 2):
        score = st.dk[bai - 1]
        vals: List = [st.stt, st.name]
        if is_num(score):
            tcs = split_score(round(score, 1), rng=rng)
            total = round(sum(tcs), 1)
            vals += [round(score, 1), total] + tcs + [total, None]
            count += 1
        else:
            vals += [DASH, DASH] + [DASH] * N_CRITERIA + [DASH, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font, c.border = font_n, border
            c.alignment = left if j == 2 else center
            if j >= 3 and isinstance(v, float):
                c.number_format = "0.0"
        ws.row_dimensions[i].height = 20

    widths = {"A": 5.66, "B": 26.66, "C": 9.66, "D": 8, "O": 8, "P": 8}
    for k in range(5, 5 + N_CRITERIA):
        widths[get_column_letter(k)] = 8
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["D"].hidden = True       # cột "Tính" phụ (width 0 trong mẫu)
    ws.freeze_panes = "C2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"
    wb.properties.title = f"Phiếu đánh giá kết quả thực hành - Bài {bai}"
    wb.save(path)
    return count
