"""
Xuất 4 biểu mẫu (DK DT, TL, LT, TK) ra file .xlsx — cùng bố cục / cột / hàng / khung kẻ như bản PDF
và như các sheet trong Form_Export (Times New Roman, in sẵn khổ giấy, tiêu đề lặp, footer).
"""
from __future__ import annotations

from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L

from . import forms_pdf as F
from .core import Meta, Student, fmt_pct, fmt_score, ngay_thang_nam, xet_dieu_kien
from .footer import FooterOptions

FONT = "Times New Roman"
_thin = Side(style="thin")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


class Sheet:
    def __init__(self, title: str, units, row_heights=None):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = title
        for i, w in enumerate(units, 1):
            self.ws.column_dimensions[L(i)].width = w
        self.r = 0  # số hàng đã dùng

    def rows(self, heights):
        """Cấp các hàng liên tiếp với chiều cao (pt); trả về số hàng (1-based) của hàng đầu."""
        first = self.r + 1
        for h in heights:
            self.r += 1
            self.ws.row_dimensions[self.r].height = h
        return first

    def cell(self, r0, r1, c0, c1, text="", size=13, bold=False, italic=False, h="center", v="center",
             border=True, wrap=True, underline=False):
        """r0,r1 (1-based, gồm cả hai đầu) ; c0,c1 (0-based)."""
        ws = self.ws
        if r0 != r1 or c0 != c1:
            ws.merge_cells(start_row=r0, end_row=r1, start_column=c0 + 1, end_column=c1 + 1)
        top = ws.cell(r0, c0 + 1)
        top.value = text if text != "" else None
        top.font = Font(name=FONT, size=size, bold=bold, italic=italic, underline="single" if underline else None)
        top.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
        if border:
            for rr in range(r0, r1 + 1):
                for cc in range(c0 + 1, c1 + 2):
                    ws.cell(rr, cc).border = _BOX

    def page(self, landscape, scale, margins, footer_margin, area, titles, ft: FooterOptions, key: str, fit_width=False):
        ws = self.ws
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        if fit_width:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        else:
            ws.page_setup.scale = scale
        l, r, t, b = margins
        ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom = l, r, t, b
        ws.page_margins.footer = footer_margin
        ws.print_area = area
        if titles:
            ws.print_title_rows = titles
        ws.print_options.horizontalCentered = True
        apply_footer(ws, ft, key)

    def save(self, path):
        self.wb.save(path)


def apply_footer(ws, ft: FooterOptions, key: str) -> None:
    left, center, right = ft.left(key), ft.center(), "Trang &P/&N"
    for part, txt in ((ws.oddFooter.left, left), (ws.oddFooter.center, center), (ws.oddFooter.right, right)):
        part.text = txt
        part.size = 12
        part.font = f"{FONT},Regular"


def _last_row(sh: Sheet) -> int:
    return sh.r


# ===========================================================================
def xlsx_dieu_kien(path: str, meta: Meta, students: List[Student], opt: F.DKOptions, ft: FooterOptions) -> None:
    sh = Sheet("DK DT", F.DK_UNITS)
    last = len(F.DK_UNITS) - 1
    ir = [33, 17, 20, 17, 25.5, 20.25, 20.25, 20.25, 20.25, 20.25, 45, 20.25, 4.5]
    r = sh.rows(ir)  # r = 1
    R = lambda i: r + i  # hàng (0-based -> 1-based)
    X = lambda b: "X" if b else "-"
    c = sh.cell
    c(R(0), R(0), 0, 6, f"PHÒNG ĐÀO TẠO\nTỔ BỘ MÔN {opt.to_bo_mon}", border=False)
    c(R(0), R(0), 7, last, "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc", border=False)
    c(R(2), R(2), 0, last, "DANH SÁCH XÉT ĐIỀU KIỆN HSSV DỰ THI MH/MĐ", size=16, bold=True, border=False)
    lf = dict(h="left", v="bottom", border=False, wrap=False)
    rt = dict(h="right", v="bottom", border=False, wrap=False)
    c(R(4), R(4), 0, 2, meta.a5_raw, **lf)
    c(R(4), R(4), 3, 4, f"SỐ TC: {meta.tin_chi}", **lf)
    c(R(4), R(4), 5, 6, "Tổng số giờ:", **rt)
    c(R(4), R(4), 7, 7, f"{meta.tong_gio}", **lf)
    c(R(5), R(5), 0, 2, f"LỚP: {meta.lop_khoa}", **lf)
    c(R(5), R(5), 6, 6, "LT:", **rt)
    c(R(5), R(5), 7, 7, f"{meta.gio_lt:g}", **lf)
    c(R(6), R(6), 0, 2, f"MÃ LHP: {opt.ma_lhp}", **lf)
    c(R(6), R(6), 6, 6, "TH:", **rt)
    c(R(6), R(6), 7, 7, f"{meta.gio_th:g}", **lf)
    c(R(7), R(7), 0, 2, "HÌNH THỨC THI (*):", **lf)
    t9, t10, t11, t12 = R(8), R(9), R(10), R(11)
    s11 = dict(size=11)
    c(t9, t9, 1, 7, "Lý Thuyết", bold=True, **s11)
    c(t9, t9, 8, last, "Thi Thực hành", bold=True, **s11)
    it = dict(italic=True, size=11)
    c(t10, t10, 1, 2, "Trắc nghiệm", **it)
    c(t10, t11, 3, 3, "Vấn đáp", **it)
    c(t10, t11, 4, 5, "Tự luận", **it)
    c(t10, t11, 6, 7, "HT khác", **it)
    c(t10, t11, 8, 10, "Địa điểm thi", **it)
    c(t10, t10, 11, 13, "Thi theo ca", **it)
    c(t10, t10, 14, last, "Thi tập trung cả lớp", **it)
    c(t11, t11, 1, 1, "Trên máy", **it)
    c(t11, t11, 2, 2, "Trên giấy", **it)
    c(t11, t11, 11, 12, "Số HSSV/ ca thi", **it)
    c(t11, t11, 13, 13, "Thời gian", **it)
    c(t11, t11, 14, last, "Thời gian thi (Phút)", **it)
    c(t12, t12, 1, 1, X(opt.tn_may), **s11)
    c(t12, t12, 2, 2, X(opt.tn_giay), **s11)
    c(t12, t12, 3, 3, X(opt.van_dap), **s11)
    c(t12, t12, 4, 5, X(opt.tu_luan), **s11)
    c(t12, t12, 6, 7, X(opt.ht_khac), **s11)
    c(t12, t12, 8, 10, opt.dia_diem or "-", **s11)
    c(t12, t12, 11, 12, opt.so_hssv_ca or "-", **s11)
    c(t12, t12, 13, 13, opt.tg_ca or "-", **s11)
    c(t12, t12, 14, last, opt.tg_tap_trung or "-", **s11)
    # tiêu đề bảng (3 hàng)
    h0 = sh.rows([57, 43.5, 43.5])
    a, b, d = h0, h0 + 1, h0 + 2
    B = dict(bold=True)
    c(a, d, 0, 0, "STT", **B)
    c(a, d, 1, 1, "Mã số HSSV", **B)
    c(a, d, 2, 2, "HỌ VÀ TÊN HSSV", **B)
    c(a, a, 3, 4, "ĐIỂM KT THƯỜNG XUYÊN", **B)
    c(b, d, 3, 3, "1", **B)
    c(b, d, 4, 4, "2", **B)
    c(a, a, 5, 10, "ĐIỂM KT ĐỊNH KỲ", **B)
    for i in range(6):
        c(b, d, 5 + i, 5 + i, str(i + 1), **B)
    c(a, d, 11, 11, "TBC KT", **B)
    c(a, b, 12, 12, "% Tiết vắng học (có phép và không phép)", **B)
    c(d, d, 12, 12, "LT + TH", **B)
    c(a, a, 13, 14, "XÉT ĐK DỰ THI\n(đánh dấu x)", **B)
    c(b, d, 13, 13, "Đủ ĐK thi", **B)
    c(b, d, 14, 14, "Không đủ ĐK thi", **B)
    c(a, d, 15, 15, "GHI CHÚ", **B)
    # dữ liệu
    rows = F.dk_rows(students, opt)
    first = sh.rows([20] * len(rows))
    for i, rw in enumerate(rows):
        for ci, txt in enumerate(rw):
            c(first + i, first + i, ci, ci, txt, h="left" if ci in (2, 15) else "center")
    # chữ ký
    f0 = sh.rows([6, 20.25, 32.25, 71.25])
    c(f0 + 1, f0 + 1, 12, last, ngay_thang_nam(opt.ngay_lap), italic=True, border=False)
    c(f0 + 2, f0 + 2, 0, 4, f"TỔ BỘ MÔN {opt.to_bo_mon}\nTỔ PHÓ", bold=True, v="top", border=False)
    c(f0 + 2, f0 + 2, 12, last, "GIÁO VIÊN GIẢNG DẠY", bold=True, v="top", border=False)
    sh.page(True, 65, (0.394, 0.0, 0.394, 0.394), 0.197, f"A1:P{sh.r}", f"{h0}:{d}", ft, "dk")
    sh.save(path)


# ===========================================================================
def xlsx_thi_lai(path: str, meta: Meta, ds: List[dict], opt: F.TLOptions, ft: FooterOptions) -> None:
    sh = Sheet("TL", F.TL_UNITS)
    ir = [15.75, 15, 26.25, 18, 26.25, 17, 17, 17, 15]
    r = sh.rows(ir)
    R = lambda i: r + i
    c = sh.cell
    lf = dict(h="left", border=False, wrap=False, size=11, bold=True)
    c(R(0), R(0), 0, 2, "PHÒNG ĐÀO TẠO", size=11, border=False)
    c(R(0), R(0), 3, 5, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", size=10, border=False)
    c(R(1), R(1), 0, 2, f"TỔ BỘ MÔN {opt.to_bo_mon}", size=11, bold=True, border=False)
    c(R(1), R(1), 3, 5, "Độc lập - Tự do - Hạnh phúc", size=11, bold=True, border=False, underline=True)
    c(R(3), R(3), 0, 5, "DANH SÁCH HSSV THI LẠI MH/MĐ (LẦN 2)", size=14, bold=True, v="bottom", border=False)
    c(R(5), R(5), 0, 2, meta.a5_raw, **lf)
    c(R(5), R(5), 3, 3, f"SỐ TC: {meta.tin_chi}", **lf)
    c(R(5), R(5), 4, 4, "Tổng số giờ:", h="right", border=False, wrap=False, size=11, bold=True)
    c(R(5), R(5), 5, 5, f"{meta.tong_gio}", **lf)
    c(R(6), R(6), 0, 2, f"LỚP: {meta.lop_khoa}", **lf)
    c(R(6), R(6), 4, 4, "LT:", h="right", border=False, wrap=False, size=11, bold=True)
    c(R(6), R(6), 5, 5, f"{meta.gio_lt:g}", **lf)
    c(R(7), R(7), 0, 2, f"MÃ LHP: {opt.ma_lhp}", **lf)
    c(R(7), R(7), 4, 4, "TH:", h="right", border=False, wrap=False, size=11, bold=True)
    c(R(7), R(7), 5, 5, f"{meta.gio_th:g}", **lf)
    h0 = sh.rows([31.5, 30])
    B = dict(bold=True, size=10)
    c(h0, h0 + 1, 0, 0, "STT", **B)
    c(h0, h0 + 1, 1, 1, "Mã số HSSV", **B)
    c(h0, h0 + 1, 2, 2, "HỌ VÀ TÊN HSSV", **B)
    c(h0, h0, 3, 4, "NỘI DUNG THI LẠI", **B)
    c(h0 + 1, h0 + 1, 3, 3, "LT", **B)
    c(h0 + 1, h0 + 1, 4, 4, "TH", **B)
    c(h0, h0 + 1, 5, 5, "GHI CHÚ", **B)
    rows = F.tl_rows(ds)
    first = sh.rows([33] * len(rows))
    al = ["center", "left", "left", "center", "center", "left"]
    for i, rw in enumerate(rows):
        for ci, txt in enumerate(rw):
            c(first + i, first + i, ci, ci, txt, h=al[ci])
    f0 = sh.rows([6, 20, 14, 60])
    c(f0 + 1, f0 + 1, 3, 5, ngay_thang_nam(opt.ngay_lap), size=11, italic=True, border=False)
    c(f0 + 2, f0 + 2, 0, 2, "TRƯỞNG BỘ MÔN", size=11, bold=True, border=False)
    c(f0 + 2, f0 + 2, 3, 5, "GIÁO VIÊN BỘ MÔN", size=11, bold=True, border=False)
    sh.page(True, 110, (0.787, 0.394, 0.591, 0.591), 0.394, f"A1:F{sh.r}", f"{h0}:{h0 + 1}", ft, "tl")
    sh.save(path)


# ===========================================================================
def xlsx_diem_lt(path: str, meta: Meta, ds: List[dict], opt: F.LTOptions, ft: FooterOptions) -> None:
    sh = Sheet("LT", F.LT_UNITS)
    ir = [15.75, 15, 15, 20.25, 18, 18.75, 18, 18, 18, 18, 19.5]
    r = sh.rows(ir)
    R = lambda i: r + i
    c = sh.cell
    c(R(0), R(0), 0, 2, "TRƯỜNG CAO ĐẲNG DẦU KHÍ", size=12, bold=True, border=False)
    c(R(0), R(0), 4, 5, "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM", size=12, border=False)
    c(R(1), R(1), 0, 2, "PHÒNG ĐÀO TẠO", size=12, bold=True, border=False)
    c(R(1), R(1), 4, 5, "Độc lập - Tự do - Hạnh phúc", size=12, bold=True, border=False, underline=True)
    c(R(2), R(2), 0, 2, f"TỔ BỘ MÔN {opt.to_bo_mon}", size=12, bold=True, border=False)
    c(R(4), R(4), 0, 5, "BẢNG ĐIỂM TỔNG HỢP THEO LỚP", size=14, bold=True, border=False)
    lf = dict(h="left", border=False, wrap=False)
    c(R(6), R(6), 0, 5, f"Tên MH/MĐ: {meta.ten}", **lf)
    c(R(7), R(7), 0, 5, f"Lần thi: {opt.lan_thi:02d}", **lf)
    c(R(8), R(8), 0, 5, f"Lớp: {meta.lop_khoa}", **lf)
    c(R(9), R(9), 0, 5, f"Ca thi: {opt.ca_thi}", **lf)
    h0 = sh.rows([33])
    for i, t in enumerate(["Stt", "Mã số HSSV", "Họ và tên", "Ngày sinh", "Điểm", "Ghi chú"]):
        c(h0, h0, i, i, t, bold=True)
    rows = F.lt_rows(ds)
    first = sh.rows([20] * len(rows))
    al = ["center", "center", "left", "center", "center", "center"]
    for i, rw in enumerate(rows):
        for ci, txt in enumerate(rw):
            c(first + i, first + i, ci, ci, txt, h=al[ci])
    f0 = sh.rows([6, 20.25, 18, 18, 50])
    c(f0 + 1, f0 + 1, 3, 5, ngay_thang_nam(opt.ngay_lap), italic=True, border=False)
    c(f0 + 2, f0 + 2, 3, 5, "CÁN BỘ GIÁM SÁT / CÁN BỘ COI THI", bold=True, border=False)
    c(f0 + 3, f0 + 3, 3, 5, "GIÁM KHẢO 1 / GIÁM KHẢO 2", bold=True, border=False)
    sh.page(False, 76, (0.591, 0.394, 0.591, 0.591), 0.315, f"A1:F{sh.r}", f"{h0}:{h0}", ft, "lt")
    sh.save(path)


# ===========================================================================
def xlsx_xep_thu_tu(path: str, meta: Meta, data: Dict[str, list], ft: FooterOptions) -> None:
    sh = Sheet("TK", F.TK_UNITS)
    cap = sh.rows([34])
    sh.cell(cap, cap, 0, len(F.TK_UNITS) - 1,
            f"XẾP THỨ TỰ ĐIỂM TỔNG KẾT – {meta.ten} – LỚP {meta.lop_khoa}", size=14, bold=True, border=False, wrap=False)
    h0 = sh.rows([18, 18])
    for c0, title, _ in F.TK_GROUPS:
        sh.cell(h0, h0, c0, c0 + 2, title, size=12, bold=True, border=False)
        sh.cell(h0 + 1, h0 + 1, c0, c0, "Stt", bold=True)
        sh.cell(h0 + 1, h0 + 1, c0 + 1, c0 + 1, "Họ và tên", bold=True)
        sh.cell(h0 + 1, h0 + 1, c0 + 2, c0 + 2, "Điểm tổng", bold=True)
    n = max([len(data[k]) for _, _, k in F.TK_GROUPS] + [0])
    first = sh.rows([20.25] * n)
    for i in range(n):
        for c0, _, key in F.TK_GROUPS:
            lst = data[key]
            if i < len(lst):
                nm, sc = lst[i]
                sh.cell(first + i, first + i, c0, c0, str(i + 1))
                sh.cell(first + i, first + i, c0 + 1, c0 + 1, nm, h="left")
                sh.cell(first + i, first + i, c0 + 2, c0 + 2, sc)
    sh.page(True, 0, (0.7, 0.7, 0.75, 0.75), 0.3, f"A1:{L(len(F.TK_UNITS))}{sh.r}", f"{h0}:{h0 + 1}", ft, "tk",
            fit_width=True)
    sh.save(path)
