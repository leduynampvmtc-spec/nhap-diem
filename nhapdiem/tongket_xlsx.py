"""
Tab "Tính điểm tổng kết": ghi kết quả tính (TBC_KT, L1, TK_L1, L2, TK_L2 của VBA) trở lại một BẢN SAO của
Form nhập điểm — các cột D, O, P, Q, R, S, T, U, V, W, X, Y — và tô màu giống macro.
"""
from __future__ import annotations

from typing import List

import openpyxl
from openpyxl.styles import Font, PatternFill

from .core import DASH, Meta, Student, is_num

YELLOW = PatternFill("solid", fgColor="FFFF00")     # 65535
GREEN = PatternFill("solid", fgColor="92D050")      # 5296274
ORANGE = PatternFill("solid", fgColor="FFC000")     # 49407
BLUE = PatternFill("solid", fgColor="00B0F0")       # 15773696
BOLD = Font(bold=True)
COLS = {"D": 4, "O": 15, "P": 16, "Q": 17, "R": 18, "S": 19, "T": 20, "U": 21, "V": 22, "W": 23, "X": 24, "Y": 25}


def _put(ws, row: int, col: str, v) -> None:
    c = ws.cell(row, COLS[col])
    c.value = v
    if is_num(v):
        c.number_format = "0.0"


def _band(c, v, lt5_yellow: bool = True) -> None:
    """Tô theo ngưỡng như macro: < 5 hoặc "-" vàng đậm; >= 8 xanh lá đậm; 7–8 cam đậm."""
    if v == DASH or (is_num(v) and v < 5):
        if lt5_yellow:
            c.fill = YELLOW
            c.font = BOLD
    elif is_num(v) and v >= 8:
        c.fill = GREEN
        c.font = BOLD
    elif is_num(v) and v >= 7:
        c.fill = ORANGE
        c.font = BOLD


def ghi_ket_qua(src: str, out: str, meta: Meta, students: List[Student], lan: int = 1) -> int:
    wb = openpyxl.load_workbook(src)
    ws = wb["DIEM TK"] if "DIEM TK" in wb.sheetnames else wb.worksheets[0]
    for st in students:
        r = st.row
        if meta.tin_chi <= 3:
            _put(ws, r, "D", DASH)
        vals = {"O": st.o, "P": st.p, "Q": st.q, "R": st.r, "S": st.s}
        if lan == 2:
            vals.update({"T": st.t, "U": st.u, "V": st.v, "W": st.w})
        else:
            vals.update({"T": DASH, "U": DASH, "V": DASH, "W": DASH})   # TK_L1 đặt T:W = "-"
        for col, v in vals.items():
            if v is not None:
                _put(ws, r, col, v)
        _put(ws, r, "X", st.x or None)
        _put(ws, r, "Y", st.y or None)
        # --- tô màu (xóa màu cũ ở các cột tính) ---
        for col in "OPQRSTUVWXY":
            c = ws.cell(r, COLS[col])
            c.fill = PatternFill(fill_type=None)
            c.font = Font(bold=False)
        for j in range(3, 16):                                 # C..O: điểm < 5 tô vàng (TBC_KT)
            c = ws.cell(r, j)
            if is_num(c.value) and c.value < 5:
                c.fill = YELLOW
        if is_num(st.o) and st.o >= 7:
            _band(ws.cell(r, COLS["O"]), st.o)
        for col, v in (("P", st.p), ("Q", st.q), ("R", st.r)):
            if v == DASH or (is_num(v) and v < 5):
                _band(ws.cell(r, COLS[col]), v)
        _band(ws.cell(r, COLS["S"]), st.s)
        if lan == 2 and is_num(st.w):
            _band(ws.cell(r, COLS["W"]), st.w)
        x = ws.cell(r, COLS["X"])
        if st.x == "Thi lại":
            x.fill = YELLOW
        elif st.x == "Học lại":
            x.fill = BLUE
    wb.save(out)
    return len(students)
