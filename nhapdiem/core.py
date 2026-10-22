"""
Lõi tính toán — port từ code VBA trong "Form Code VBA.xlsm" sang Python thuần.

Nguồn VBA tương ứng:
  TBC_KT      -> compute_tbc_kt()      (cột O  - TBC KT)
  L1          -> compute_lan1()        (cột R  - điểm TBC thi lần 1)
  TK_L1       -> compute_lan1()        (cột S  - điểm tổng kết, X - đánh giá, Y - ghi chú)
  Tinhnghihoc -> load_attendance()     (% tiết vắng)
  Xet_DK_2021 / Xet_DK_2022 -> xet_dieu_kien()
  THIL        -> danh_sach_thi_lai()
  LYT / LYT2  -> bang_diem_lt()
  SQL_lon7 / SQL_nho5 / SQL_3lonnhat / SQL_3nhonhat -> xep_thu_tu()

Quy ước giá trị ô (giống Excel):
  - số      -> float
  - "-"     -> DASH  (ô gạch ngang: không có điểm / học lại)
  - để trống-> None  (Empty trong VBA)
Không dùng FileSystemObject / Scripting.Dictionary / ADODB / Excel COM nên chạy
giống hệt nhau trên Windows và macOS.
"""
from __future__ import annotations

import datetime as _dt
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
from typing import Dict, List, Optional, Tuple, Union

import openpyxl

DASH = "-"
Cell = Union[float, str, None]

SHEET_NAME = "DIEM TK"
FIRST_ROW = 10  # hàng đầu tiên chứa HSSV trong sheet DIEM TK

# Cột (1-based) trong sheet DIEM TK
COL_NAME = 2      # B  Họ và tên
COL_TX1 = 3       # C  KT thường xuyên 1
COL_TX2 = 4       # D  KT thường xuyên 2
COL_DK1 = 5       # E  KT định kỳ 1 ... N = 10
COL_DK_LAST = 14  # N
COL_THI_LT1 = 16  # P  điểm thi lần 1 - LT
COL_THI_TH1 = 17  # Q  điểm thi lần 1 - TH
COL_THI_LT2 = 20  # T  điểm thi lần 2 - LT
COL_THI_TH2 = 21  # U  điểm thi lần 2 - TH
COL_MSSV = 27     # AA
COL_DOB = 31      # AE
COL_DE = 32       # AF


# ----------------------------------------------------------------------------
# Tiện ích số / chuỗi
# ----------------------------------------------------------------------------
def xl_round(x: float, nd: int = 0) -> float:
    """ROUND của Excel (làm tròn nửa ra xa số 0) — Application.Round trong VBA."""
    q = Decimal(1).scaleb(-nd)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def bo_dau(text: str) -> str:
    """Bỏ dấu tiếng Việt (tương đương hàm bo_dau trong VBA, giữ nguyên chữ hoa/thường)."""
    if text is None:
        return ""
    s = str(text).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return unicodedata.normalize("NFC", s)


def nfc_upper(text: str) -> str:
    return unicodedata.normalize("NFC", str(text or "")).strip().upper()


def is_num(v: Cell) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def norm_cell(v) -> Cell:
    """Chuẩn hóa giá trị đọc từ Excel về float / DASH / None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return None
    s = str(v).strip()
    if s == "":
        return None
    if s in ("-", "–", "—"):
        return DASH
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return DASH  # chuỗi lạ (vd "V", "KP") coi như không có điểm


def fmt_score(v: Cell, nd: int = 1) -> str:
    """Hiển thị điểm theo định dạng 0.0 (giống NumberFormat="0.0")."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return f"{v:.{nd}f}"


def fmt_pct(v: Optional[float]) -> str:
    """Định dạng 0% ; giá trị 0 / None hiển thị rỗng (VBA: If M = 0 Then M = "")."""
    if v is None or v == 0:
        return ""
    return f"{round(v * 100):d}%"


def fmt_date(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, _dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def ngay_thang_nam(d: _dt.date) -> str:
    return f"T.p Hồ Chí Minh, ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


# ----------------------------------------------------------------------------
# Mô hình dữ liệu
# ----------------------------------------------------------------------------
@dataclass
class Meta:
    a5_raw: str = ""          # "TÊN MÔN HỌC: CƠ KỸ THUẬT"  (A5 nguyên văn)
    loai: str = ""            # "MÔN HỌC" | "MÔ ĐUN"
    ten: str = ""             # "CƠ KỸ THUẬT"
    tong_gio: int = 0
    gio_lt: float = 0.0
    gio_th: float = 0.0
    tin_chi: int = 0
    lop: str = ""             # "SCTBCBDK 1"
    khoa: str = ""            # "2025"
    he: str = ""

    @property
    def lop_khoa(self) -> str:        # AG4: SCTBCBDK 1-2025
        return f"{self.lop}-{self.khoa}" if self.khoa else self.lop

    @property
    def ten_khong_dau(self) -> str:   # AE6: bo_dau(tên môn)
        return bo_dau(self.ten).strip()

    @property
    def la_mo_dun(self) -> bool:
        return nfc_upper(self.loai) == nfc_upper("MÔ ĐUN")

    @property
    def la_mon_hoc(self) -> bool:
        return nfc_upper(self.loai) == nfc_upper("MÔN HỌC")


@dataclass
class Student:
    row: int
    stt: int
    name: str
    tx1: Cell = None
    tx2: Cell = None
    dk: List[Cell] = field(default_factory=lambda: [None] * 10)
    mssv: str = ""
    dob: object = None
    de: object = None
    # đọc từ file
    p: Cell = None   # thi lần 1 LT
    q: Cell = None   # thi lần 1 TH
    t: Cell = None   # thi lần 2 LT
    u: Cell = None   # thi lần 2 TH
    # tính toán
    o: Cell = None   # TBC KT
    r: Cell = None   # TBC thi lần 1
    s: Cell = None   # tổng kết lần 1
    x: str = ""      # đánh giá
    y: str = ""      # ghi chú
    # % tiết vắng (từ file điểm danh)
    vang_tong: Optional[float] = None
    vang_lt: Optional[float] = None
    vang_th: Optional[float] = None

    _tbc_raw: Optional[float] = None  # TBC chưa làm tròn (dùng tính điểm tổng kết)


# ----------------------------------------------------------------------------
# Đọc file nhập điểm
# ----------------------------------------------------------------------------
def _cell_value(ws, ws_v, row: int, col: int):
    raw = ws.cell(row, col).value
    if isinstance(raw, str) and raw.startswith("="):
        return ws_v.cell(row, col).value  # giá trị đã tính sẵn (nếu có)
    if raw is not None and hasattr(raw, "text"):  # ArrayFormula
        return ws_v.cell(row, col).value
    return raw


def read_input(path: str) -> Tuple[Meta, List[Student], List[str]]:
    """Đọc sheet DIEM TK của "Form nhập điểm". Trả về (meta, ds HSSV, cảnh báo)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    name = SHEET_NAME if SHEET_NAME in wb.sheetnames else wb.sheetnames[0]
    ws, ws_v = wb[name], wb_v[name]
    warnings: List[str] = []

    def txt(addr: str) -> str:
        v = _cell_value(ws, ws_v, ws[addr].row, ws[addr].column)
        return "" if v is None else str(v).strip()

    meta = Meta()
    meta.a5_raw = txt("A5")
    m = re.match(r"^\s*T[ÊE]N\s+(.*?)\s*:\s*(.*)$", unicodedata.normalize("NFC", meta.a5_raw), re.I)
    if m:
        meta.loai, meta.ten = m.group(1).strip(), m.group(2).strip()
    elif ":" in meta.a5_raw:
        meta.ten = meta.a5_raw.split(":", 1)[1].strip()
        meta.loai = txt("AA6")
    else:
        meta.ten = meta.a5_raw
        meta.loai = txt("AA6")
    if not meta.loai:
        meta.loai = txt("AA6")

    # SỐ GIỜ / (LT: 30; TH: 15)
    mh = re.search(r"(\d+)", txt("O5"))
    meta.tong_gio = int(mh.group(1)) if mh else 0
    q5 = txt("Q5")
    mlt = re.search(r"LT\s*:\s*(\d+(?:[.,]\d+)?)", q5, re.I)
    mth = re.search(r"TH\s*:\s*(\d+(?:[.,]\d+)?)", q5, re.I)
    if mlt and mth:
        meta.gio_lt = float(mlt.group(1).replace(",", "."))
        meta.gio_th = float(mth.group(1).replace(",", "."))
    else:
        lt = _cell_value(ws, ws_v, 5, 31)   # AE5
        th = _cell_value(ws, ws_v, 5, 32)   # AF5
        try:
            meta.gio_lt, meta.gio_th = float(lt), float(th)
        except (TypeError, ValueError):
            meta.gio_lt, meta.gio_th = float(meta.tong_gio), 0.0
            warnings.append("Không đọc được (LT: ..; TH: ..) ở ô Q5 — coi như toàn bộ là giờ LT.")
    meta.tin_chi = int(xl_round(meta.gio_lt / 15 + meta.gio_th / 30, 0))

    meta.lop = txt("B6")
    ek = txt("E6")
    meta.khoa = ek.split(":", 1)[1].strip() if ":" in ek else ek
    meta.he = txt("O6")

    students: List[Student] = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        nm = _cell_value(ws, ws_v, r, COL_NAME)
        if nm is None or str(nm).strip() == "":
            continue
        st = Student(row=r, stt=len(students) + 1, name=str(nm).strip())
        st.tx1 = norm_cell(_cell_value(ws, ws_v, r, COL_TX1))
        st.tx2 = norm_cell(_cell_value(ws, ws_v, r, COL_TX2))
        st.dk = [norm_cell(_cell_value(ws, ws_v, r, COL_DK1 + i)) for i in range(10)]
        st.p = norm_cell(_cell_value(ws, ws_v, r, COL_THI_LT1))
        st.q = norm_cell(_cell_value(ws, ws_v, r, COL_THI_TH1))
        st.t = norm_cell(_cell_value(ws, ws_v, r, COL_THI_LT2))
        st.u = norm_cell(_cell_value(ws, ws_v, r, COL_THI_TH2))
        mssv = _cell_value(ws, ws_v, r, COL_MSSV)
        st.mssv = "" if mssv is None else str(mssv).strip()
        st.dob = _cell_value(ws, ws_v, r, COL_DOB)
        st.de = _cell_value(ws, ws_v, r, COL_DE)
        students.append(st)

    if not students:
        warnings.append("Không tìm thấy HSSV nào (cột B từ hàng 10 trở xuống đang trống).")
    elif not any(s.mssv for s in students):
        warnings.append("Cột AA (Mã số HSSV) trống — cột Mã số HSSV trong file xuất sẽ để trống.")
    return meta, students, warnings


# ----------------------------------------------------------------------------
# Tính điểm  (TBC_KT, L1, TK_L1)
# ----------------------------------------------------------------------------
def _tbc_kt(st: Student, tin_chi: int) -> Tuple[Cell, Optional[float]]:
    """Trả về (O làm tròn 1 chữ số, TBC chưa làm tròn)."""
    tx2 = DASH if tin_chi <= 3 else st.tx2          # TBC_KT: nếu <= 3 tín chỉ thì D = "-"
    tx = [v for v in (st.tx1, tx2) if is_num(v)]
    dk = [v for v in st.dk if is_num(v)]
    den = len(tx) + 2 * len(dk)
    if den == 0:
        return DASH, None
    raw = (sum(tx) + 2 * sum(dk)) / den
    return xl_round(raw, 1), raw


def compute_tbc_kt(students: List[Student], meta: Meta) -> None:
    for st in students:
        if meta.tin_chi <= 3:
            st.tx2 = DASH
        st.o, st._tbc_raw = _tbc_kt(st, meta.tin_chi)
        if not is_num(st.o) or st.o < 5:             # O < 5 hoặc "-" => học lại
            st.p = DASH
            st.q = DASH
            st.x = "Học lại"
        else:
            st.x = ""


def _lt5(v: Cell, n: float) -> bool:
    """So sánh kiểu VBA Variant: ô trống = 0; "-" (chuỗi) luôn > số."""
    if v is None:
        return 0 < n
    if is_num(v):
        return v < n
    return False


def compute_lan1(students: List[Student], meta: Meta) -> None:
    lt_h, th_h = meta.gio_lt, meta.gio_th
    w_lt, w_th = lt_h / 15, th_h / 30
    for st in students:
        o = st.o
        o_ok = is_num(o) and o >= 5
        # ---- L1: điểm TBC thi lần 1 (R) ----
        r: Cell = None
        if meta.la_mo_dun:
            if o_ok:
                if st.q != DASH:
                    if st.p != DASH:
                        if is_num(st.p) and is_num(st.q) and (w_lt + w_th) > 0:
                            r = xl_round((st.p * w_lt + st.q * w_th) / (w_lt + w_th), 1)
                    else:
                        r = st.q
        elif meta.la_mon_hoc:
            if o_ok:
                st.q = DASH
                r = st.p
        if not o_ok:
            st.p = st.q = r = DASH
        st.r = r

        # ---- TK_L1: điểm tổng kết (S) ----
        st.x, st.y = "", ""
        tbc = st._tbc_raw
        s: Cell = None
        if meta.la_mo_dun:
            if o_ok and tbc is not None:
                if st.p != DASH and is_num(st.p) and is_num(st.q) and (w_lt + w_th) > 0:
                    s = xl_round(((st.p * w_lt + st.q * w_th) / (w_lt + w_th)) * 0.6 + tbc * 0.4, 1)
                elif st.p == DASH and is_num(r):
                    s = xl_round(r * 0.6 + tbc * 0.4, 1)
                elif st.p == DASH and r is None:
                    s = xl_round(0 * 0.6 + tbc * 0.4, 1)
            else:
                s = DASH
        elif meta.la_mon_hoc:
            if o_ok and tbc is not None:
                rv = r if is_num(r) else 0.0
                s = xl_round(rv * 0.6 + tbc * 0.4, 1)
            else:
                s = DASH
        st.s = s

        notice = "Điểm thi lần thứ nhất < 5"
        P, Q, R, O, S = st.p, st.q, st.r, st.o, st.s
        all_ok = (is_num(S) and S >= 5 and is_num(O) and O >= 5 and is_num(R) and R >= 5)
        if all_ok:
            st.x = "Đạt"
        elif _lt5(R, 5) and R != DASH and o_ok and th_h == 0:
            st.x, st.y = "Thi lại", notice
        elif P != DASH and _lt5(P, 3):
            st.x, st.y = "Thi lại", notice
        elif Q != DASH and _lt5(Q, 3):
            st.x, st.y = "Thi lại", notice
        elif is_num(P) and P >= 3 and R != DASH and _lt5(R, 5):
            st.x, st.y = "Thi lại", notice
        elif is_num(Q) and Q >= 3 and R != DASH and _lt5(R, 5):
            st.x, st.y = "Thi lại", notice
        elif R == DASH:
            st.x, st.y = "Học lại", "Nghỉ học quá số buổi quy định"
        if th_h == 0:
            st.q = DASH


def compute_all(students: List[Student], meta: Meta) -> None:
    compute_tbc_kt(students, meta)
    compute_lan1(students, meta)


# ----------------------------------------------------------------------------
# Tiết vắng (Tinhnghihoc)
# ----------------------------------------------------------------------------
def _read_att_sheet(ws) -> Dict[str, Tuple[float, float, float]]:
    """
    Port macro Tinhnghihoc (Xet_DK.bas). Bố cục file điểm danh:
      - hàng 7 = 1 (mỗi buổi trọng số 1), hàng 10 = "LT"/"TH", hàng 11 = thời lượng buổi (có số = 1 buổi)
      - hàng 12+: cột B = họ tên; ô buổi = "X" / "P" / "KP"  (P và KP đều là vắng)
    % vắng LT = số buổi LT vắng / số buổi (ô có số ở hàng 11); % vắng TH tương tự; tổng = LT + TH.
    Mẫu số là TỔNG số buổi theo thời khoá biểu (kể cả buổi chưa diễn ra) — đúng như VBA.
    """
    cols = []
    for c in range(4, ws.max_column + 1):
        kind = str(ws.cell(10, c).value or "").strip().upper()
        h = ws.cell(11, c).value
        if kind in ("LT", "TH") and isinstance(h, (int, float)) and not isinstance(h, bool):
            cols.append((c, kind))
    if not cols:
        raise ValueError("Không tìm thấy các buổi học (hàng 10 = LT/TH, hàng 11 = số giờ) trong sheet điểm danh.")
    n_sess = len(cols)
    out: Dict[str, Tuple[float, float, float]] = {}
    for r in range(12, ws.max_row + 1):
        nm = ws.cell(r, 2).value
        if nm in (None, "") or not str(nm).strip():
            continue
        stt = ws.cell(r, 1).value
        if not isinstance(stt, (int, float)) and not str(stt or "").strip().isdigit():
            continue          # hàng ghi chú/chân trang (không có STT)
        a = {"LT": 0, "TH": 0}
        for c, k in cols:
            if str(ws.cell(r, c).value or "").strip().upper() in ("P", "KP"):
                a[k] += 1
        key = nfc_upper(bo_dau(" ".join(str(nm).split())))
        if key not in out:
            out[key] = ((a["LT"] + a["TH"]) / n_sess, a["LT"] / n_sess, a["TH"] / n_sess)
    return out


def attendance_sheets(path: str) -> List[str]:
    return openpyxl.load_workbook(path, read_only=False, data_only=True).sheetnames


def load_attendance(path: str, sheet: Optional[str] = None,
                    names: Optional[List[str]] = None) -> Dict[str, Tuple[float, float, float]]:
    """
    Đọc file điểm danh. Nếu không chỉ định `sheet`, chọn sheet khớp nhiều tên HSSV nhất với `names`
    (danh sách trong file nhập điểm); nếu không có `names` lấy sheet đầu tiên đọc được.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            raise ValueError(f"File điểm danh không có sheet '{sheet}'.")
        return _read_att_sheet(wb[sheet])
    best, best_n, err = None, -1, None
    keys = {nfc_upper(bo_dau(" ".join(n.split()))) for n in (names or [])}
    for ws in wb.worksheets:
        try:
            d = _read_att_sheet(ws)
        except ValueError as e:
            err = e
            continue
        n = len(keys & set(d))
        if n > best_n:
            best, best_n = d, n
    if best is None:
        raise err or ValueError("File điểm danh không có sheet hợp lệ.")
    return best


def apply_attendance(students: List[Student], att: Dict[str, Tuple[float, float, float]]) -> List[str]:
    """Gán % vắng cho HSSV theo tên không dấu. Trả về danh sách tên không khớp."""
    missing = []
    for st in students:
        key = nfc_upper(bo_dau(" ".join(st.name.split())))
        if key in att:
            st.vang_tong, st.vang_lt, st.vang_th = att[key]
        else:
            st.vang_tong = st.vang_lt = st.vang_th = None
            missing.append(st.name)
    return missing


# ----------------------------------------------------------------------------
# Điều kiện dự thi (Xet_DK_2021 / Xet_DK_2022)
# ----------------------------------------------------------------------------
HOC_LAI = "Học lại"


def xet_dieu_kien(st: Student, quy_che: str = "2022") -> Tuple[bool, str]:
    """Trả về (đủ ĐK?, ghi chú)."""
    L = st.o
    if quy_che == "2021":
        # vắng LT > 30%  hoặc  điểm TBC KT < 5  hoặc  vắng TH > 0
        if (st.vang_lt or 0) > 0.3 or (is_num(L) and L < 5) or (st.vang_th or 0) > 0:
            return False, HOC_LAI
        return True, ""
    if (st.vang_tong or 0) > 0.2 or (is_num(L) and L < 5):
        return False, HOC_LAI
    if L == DASH or L is None:
        return False, HOC_LAI
    return True, ""


# ----------------------------------------------------------------------------
# Thi lại (THIL)
# ----------------------------------------------------------------------------
def danh_sach_thi_lai(students: List[Student], meta: Meta, text_lt: str, text_th: str):
    """
    Trả về list dict: stt, mssv, name, lt(bool), th(bool), ghichu — đúng thứ tự VBA:
    chỉ HSSV có đánh giá "Thi lại", sắp xếp theo họ tên GIẢM DẦN (Order by F2 Desc).
    """
    sel = [s for s in students if s.x.startswith("Thi")]
    # VBA: So sánh chuỗi nhị phân, chèn ổn định, giảm dần theo họ tên
    sel = sorted(sel, key=lambda s: unicodedata.normalize("NFC", s.name), reverse=True)
    both_txt = f"{text_lt} và {text_th}"
    rows = []
    for i, s in enumerate(sel, 1):
        lt = th = False
        note = ""
        G, H = s.p, s.q
        if meta.gio_th > 0 and meta.gio_th > meta.gio_lt and meta.la_mo_dun:
            if is_num(G) and G < 5 and is_num(H) and H < 5:
                lt = th = True
                note = both_txt
            if is_num(G) and G < 5 and note != both_txt:
                lt, note = True, text_lt
            if is_num(H) and H < 5 and note != both_txt:
                th, note = True, text_th
        else:
            if (is_num(G) and G < 5) or G is None:
                lt, note = True, text_lt
        rows.append(dict(stt=i, mssv=s.mssv, name=s.name, lt=lt, th=th, ghichu=note))
    return rows


# ----------------------------------------------------------------------------
# Bảng điểm thi LT (LYT / LYT2)
# ----------------------------------------------------------------------------
def bang_diem_lt(students: List[Student], lan_thi: int = 1):
    rows = []
    if lan_thi == 1:
        sel = [(s, s.p) for s in students]
    else:
        # LYT2: HSSV có điểm thi LT lần 1 < 5 -> hiển thị điểm thi lần 2
        sel = [(s, s.t) for s in students if is_num(s.p) and s.p < 5]
    for i, (s, score) in enumerate(sel, 1):
        rows.append(dict(stt=i, mssv=s.mssv, name=s.name, dob=fmt_date(s.dob), diem=fmt_score(score), ghichu=""))
    return rows


# ----------------------------------------------------------------------------
# Xếp thứ tự (SQL_lon7 / SQL_nho5 / SQL_3lonnhat / SQL_3nhonhat)
# ----------------------------------------------------------------------------
def xep_thu_tu(students: List[Student]):
    def sc(s):
        return s.s if is_num(s.s) else None

    nums = [s for s in students if sc(s) is not None]
    lon7 = sorted([s for s in nums if sc(s) >= 7.0], key=lambda s: sc(s), reverse=True)
    nho5 = [s for s in students if sc(s) is None]                       # không có điểm -> "-"
    nho5_scored = sorted([s for s in nums if sc(s) < 5.0], key=lambda s: sc(s), reverse=True)
    top3 = sorted(nums, key=lambda s: sc(s), reverse=True)[:3]
    bottom_scored = sorted(nums, key=lambda s: sc(s))[:3]

    def pack(lst, with_dash=False):
        out = []
        for s in lst:
            out.append((s.name, DASH if sc(s) is None else fmt_score(sc(s))))
        return out

    return dict(
        lon7=pack(lon7),
        nho5=pack(nho5) + pack(nho5_scored),
        top3=pack(top3),
        bottom3=pack(nho5) + pack(bottom_scored),
    )


# ----------------------------------------------------------------------------
# Tên file xuất
# ----------------------------------------------------------------------------
_BAD = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def safe_name(s: str) -> str:
    return _BAD.sub("_", s).strip().rstrip(".")


def out_filename(prefix: str, meta: Meta, ngay: _dt.date, ext: str, extra: str = "") -> str:
    """ddmmyy_<prefix>_<TÊN MÔN KHÔNG DẤU>_<LỚP-KHÓA>[_extra].ext   (giống VBA)."""
    parts = [ngay.strftime("%d%m%y"), prefix, meta.ten_khong_dau, meta.lop_khoa]
    if extra:
        parts.append(extra)
    return safe_name("_".join(p for p in parts if p)) + "." + ext
